"""Spreadsheet / CSV export of a run — shared by the live pages and the run library.

Pure functions (no dialogs) so the file formats can be tested directly; the
Qt side lives in ``app.ui.widgets.export_ui``.

Formats
-------
* **CSV** — the readings only: one header row, then one row per reading, empty
  cell for a missing value. Kept free of comment lines so it opens unchanged in
  Origin, Excel or ``pandas.read_csv``.
* **Excel (.xlsx)** — sheet ``Data`` (the same table as the CSV, then any summary
  lines and the charts), optional extra tables (e.g. Seebeck by T₀ bin), and a
  ``Run`` sheet with who/what/when so a file on its own still says where it
  came from.
"""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from datetime import datetime
from typing import Any, List, Optional, Sequence, Tuple

Row = Sequence[Any]


@dataclass
class Sheet:
    name: str
    headers: List[str]
    rows: List[Row]


@dataclass
class RunTable:
    """What one export contains."""

    headers: List[str]
    rows: List[Row]
    #: (label, value) lines written under the data on the Data sheet.
    summary: List[Tuple[str, Any]] = field(default_factory=list)
    #: (label, value) lines for the Run sheet: run id, sample, operator, …
    info: List[Tuple[str, Any]] = field(default_factory=list)
    #: Further tables, each on its own sheet.
    extra: List[Sheet] = field(default_factory=list)


# ── file names ───────────────────────────────────────────────────────────────
_UNSAFE = re.compile(r"[^A-Za-z0-9._-]+")


def _slug(text: str, limit: int = 40) -> str:
    return _UNSAFE.sub("-", text.strip()).strip("-.")[:limit]


def default_filename(kind: str, *, run_id: Optional[int] = None, sample: Optional[str] = None,
                     what: str = "data", ext: str = "xlsx",
                     when: Optional[datetime] = None) -> str:
    """``seebeck_142_Ti3C2-04_data.xlsx`` — or, before the run has an id,
    ``seebeck_Ti3C2-04_20261020-1408_data.xlsx``."""
    parts = [_slug(kind) or "run"]
    if run_id is not None:
        parts.append(str(run_id))
    if sample and _slug(sample):
        parts.append(_slug(sample))
    if run_id is None:
        parts.append((when or datetime.now()).strftime("%Y%m%d-%H%M"))
    if what:
        parts.append(_slug(what))
    return "_".join(parts) + (f".{ext}" if ext else "")


def with_extension(path: str, selected_filter: str = "") -> Tuple[str, str]:
    """Resolve the format from the chosen filter or the typed extension and make
    sure the path ends with it. Returns ``(path, "csv" | "xlsx")``."""
    low = path.lower()
    if low.endswith(".csv") or (selected_filter.upper().startswith("CSV")
                                and not low.endswith(".xlsx")):
        return (path if low.endswith(".csv") else path + ".csv"), "csv"
    return (path if low.endswith(".xlsx") else path + ".xlsx"), "xlsx"


# ── writers ──────────────────────────────────────────────────────────────────
def _cell(v: Any) -> Any:
    return "" if v is None else v


def write_csv(path: str, table: RunTable) -> None:
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(table.headers)
        for row in table.rows:
            w.writerow([_cell(v) for v in row])


def write_xlsx(path: str, table: RunTable,
               images: Sequence[Tuple[str, bytes]] = ()) -> bool:
    """Write the workbook; ``images`` are (title, PNG bytes) placed under the
    data. Returns False when the charts could not be embedded (no Pillow) —
    the data is saved either way."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    bold = Font(bold=True)
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Data"
    ws.append(table.headers)
    for c in ws[1]:
        c.font = bold
    ws.freeze_panes = "A2"
    for row in table.rows:
        ws.append([_cell(v) for v in row])
    from openpyxl.utils import get_column_letter
    for i, h in enumerate(table.headers, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(12, len(str(h)) + 2)
    if table.summary:
        ws.append([])
        for label, value in table.summary:
            ws.append([label, _cell(value)])
            ws.cell(ws.max_row, 1).font = bold

    for sheet in table.extra:
        xs = wb.create_sheet(sheet.name[:31])
        xs.append(sheet.headers)
        for c in xs[1]:
            c.font = bold
        for row in sheet.rows:
            xs.append([_cell(v) for v in row])

    if table.info:
        rs = wb.create_sheet("Run")
        for label, value in table.info:
            rs.append([label, _cell(value)])
            rs.cell(rs.max_row, 1).font = bold
        rs.column_dimensions["A"].width = 22
        rs.column_dimensions["B"].width = 48

    embedded = True
    if images:
        try:
            from openpyxl.drawing.image import Image as XLImage
            row = ws.max_row + 3
            for title, png in images:
                ws.cell(row, 1, title).font = bold
                img = XLImage(io.BytesIO(png))
                ws.add_image(img, f"A{row + 1}")
                row += 2 + int(img.height / 20) + 1     # ~20 px per default row
        except ImportError:                              # Pillow missing
            embedded = False
    wb.save(path)
    return embedded


# ── the two measurement types ───────────────────────────────────────────────
def seebeck_table(data: Sequence[dict], columns: Sequence[Tuple[str, str]],
                  bins: Sequence[dict] = (),
                  info: Sequence[Tuple[str, Any]] = ()) -> RunTable:
    """Readings (``columns`` = (row key, header)) plus the per-T₀-bin fits."""
    t = RunTable(headers=[h for _, h in columns],
                 rows=[[r.get(k) for k, _ in columns] for r in data],
                 info=list(info))
    if bins:
        t.extra.append(Sheet(
            "Seebeck by T0 bin",
            ["T0 centre [K]", "T0 min [K]", "T0 max [K]", "S [µV/K]", "± [µV/K]",
             "Points used", "Spikes removed", "Fit note"],
            [[b.get("T0_center_K"), b.get("T0_min_K"), b.get("T0_max_K"),
              b.get("S_uV_per_K"), b.get("S_uncertainty_uV_per_K"), b.get("n_points"),
              b.get("n_rejected"), b.get("fit_warning")] for b in bins]))
        s = [b["S_uV_per_K"] for b in bins if b.get("S_uV_per_K") is not None]
        if s:
            t.summary.append(("S, mean of bins [µV/K]", round(sum(s) / len(s), 3)))
    return t


IV_HEADERS = ["Current (A)", "Voltage (V)", "Resistance (Ω)", "Power (mW)",
              "Resistivity (Ω·cm)", "Conductivity (S/cm)"]


def iv_rows(points: Sequence[dict]) -> List[List[Any]]:
    out = []
    for p in points:
        i, v = p.get("current"), p.get("voltage")
        pw = i * v * 1000 if (i is not None and v is not None) else None
        out.append([i, v, p.get("resistance"), pw, p.get("resistivity"), p.get("conductivity")])
    return out


def iv_table(points: Sequence[dict], *, fit_R: Optional[float] = None,
             fit_R2: Optional[float] = None, ohmic: Optional[str] = None,
             temperature_C: Optional[float] = None, cf: Optional[float] = None,
             rho_notes: Sequence[str] = (),
             info: Sequence[Tuple[str, Any]] = ()) -> RunTable:
    summary: List[Tuple[str, Any]] = [("Fit R (Ω)", fit_R), ("R²", fit_R2),
                                      ("Ohmic check", ohmic)]
    if temperature_C is not None:
        summary.append(("Sample temperature (°C)", temperature_C))
    if cf is not None:
        summary.append(("4PP correction factor CF", cf))
    summary += [("ρ note", n) for n in rho_notes]
    return RunTable(headers=list(IV_HEADERS), rows=list(iv_rows(points)), summary=summary,
                    info=list(info))


def iv_info(params: dict) -> List[Tuple[str, Any]]:
    """Run-sheet lines describing an I-V sweep, from its stored parameters."""
    geo = {"bar": "Bar", "vdp": "van der Pauw", "4pp": "4-point probe"}
    lines: List[Tuple[str, Any]] = []
    if params.get("geometry"):
        lines.append(("Geometry", geo.get(params["geometry"], params["geometry"])))
    for key, label in (("length", "Length (cm)"), ("width", "Width (cm)"),
                       ("thickness", "Thickness (cm)"), ("spacing", "Probe spacing (cm)")):
        if params.get(key):
            lines.append((label, params[key]))
    if params.get("source_mode"):
        unit = "A" if params["source_mode"] == "current" else "V"
        lines.append(("Sweep", f"source {params['source_mode']} {params.get('start')} → "
                               f"{params.get('stop')} {unit}, {params.get('points')} points"
                               + (", forward + reverse" if params.get("bidirectional") else "")))
    if "four_wire" in params:
        lines.append(("Wiring", "4-wire (remote sense)" if params["four_wire"] else "2-wire"))
    if params.get("notes"):
        lines.append(("Notes", params["notes"]))
    return lines


def info_lines(*, kind: str, run_id: Optional[int] = None, sample: Optional[str] = None,
               operator: Optional[str] = None, started: Optional[datetime] = None,
               finished: Optional[datetime] = None, status: Optional[str] = None,
               extra: Sequence[Tuple[str, Any]] = ()) -> List[Tuple[str, Any]]:
    """The Run sheet: identity of the run plus when/what exported it."""
    from ..version import APP_VERSION

    def ts(d: Optional[datetime]) -> str:
        return d.strftime("%Y-%m-%d %H:%M:%S") if d else ""

    lines: List[Tuple[str, Any]] = [
        ("Measurement", kind),
        ("Run", f"#{run_id}" if run_id is not None else "not saved yet"),
        ("Sample", sample or ""),
        ("Operator", operator or ""),
    ]
    if started:
        lines.append(("Started", ts(started)))
    if finished:
        lines.append(("Finished", ts(finished)))
    if status:
        lines.append(("Status", status))
    lines.extend(extra)
    lines += [("Exported", ts(datetime.now())), ("Software", f"TE Measurement {APP_VERSION}")]
    return lines
